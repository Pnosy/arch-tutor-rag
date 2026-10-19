# -*- coding: utf-8 -*-
"""检查 PPT 中所有可能溢出或错位的元素，并打印图表坐标轴设置。"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
p = ROOT / "deliverables" / "Arch-Tutor_系统展示.pptx"
prs = Presentation(str(p))
SW, SH = prs.slide_width, prs.slide_height
print(f"幻灯片尺寸: {SW/914400:.2f} x {SH/914400:.2f} in, 共 {len(prs.slides)} 页\n")

issues = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.left is None:
            continue
        r = sh.left + (sh.width or 0)
        b = sh.top + (sh.height or 0)
        if r > SW + Emu(9525) or b > SH + Emu(9525) or sh.left < 0 or sh.top < 0:
            issues += 1
            print(f"[越界] 第{i}页 shape={sh.shape_id} "
                  f"l={sh.left/914400:.2f} t={sh.top/914400:.2f} "
                  f"r={r/914400:.2f} b={b/914400:.2f}")
        if sh.has_chart:
            ch = sh.chart
            va = ch.value_axis
            print(f"[图表] 第{i}页 type={ch.chart_type} "
                  f"min={va.minimum_scale} max={va.maximum_scale}")
print(f"\n越界元素数: {issues}")
