# -*- coding: utf-8 -*-
"""导出 PPT 每页文本，用于人工核对内容与版面条目。"""
import sys
from pathlib import Path

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
prs = Presentation(str(ROOT / "deliverables" / "Arch-Tutor_实施过程.pptx"))

for i, s in enumerate(prs.slides, 1):
    print("=" * 76)
    print(f"第 {i} 页")
    print("=" * 76)
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            for line in sh.text_frame.text.split("\n"):
                if line.strip():
                    print(f"  {line}")
        if sh.has_table:
            for r in sh.table.rows:
                cells = [c.text.replace("\n", " / ") for c in r.cells]
                print("  | " + " | ".join(cells))
