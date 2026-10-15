# -*- coding: utf-8 -*-
"""校验交付物：检查 PPT 页数/溢出风险，DOCX 段落与表格数。"""
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
D = ROOT / "deliverables"

print("=" * 72)
print("交付物校验")
print("=" * 72)

# ---- PPT ----
p = D / "Arch-Tutor_实施过程.pptx"
prs = Presentation(str(p))
SW, SH = prs.slide_width, prs.slide_height
print(f"\nPPT: {p.name}")
print(f"  尺寸: {SW/914400:.2f} x {SH/914400:.2f} inch   页数: {len(prs.slides)}")
issues = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        # 越界检查
        if sh.left is None:
            continue
        r = sh.left + (sh.width or 0)
        b = sh.top + (sh.height or 0)
        if r > SW + Emu(10000) or b > SH + Emu(10000):
            issues += 1
            print(f"  [溢出] 第{i}页 {sh.shape_type} right={r/914400:.2f} bottom={b/914400:.2f}")
    # 统计文本量，评估拥挤
    chars = 0
    for sh in s.shapes:
        if sh.has_text_frame:
            chars += len(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    chars += len(c.text)
    flag = "  <-- 文字较多" if chars > 1100 else ""
    print(f"  第{i:>2}页  形状={len(s.shapes):>2}  字符={chars:>4}{flag}")
print(f"\n  越界形状数: {issues}")

# ---- DOCX ----
for name in ("Arch-Tutor_测试报告.docx", "Arch-Tutor_小组分工说明.docx"):
    f = D / name
    doc = Document(str(f))
    nchars = sum(len(x.text) for x in doc.paragraphs)
    ntab = len(doc.tables)
    tchars = sum(len(c.text) for t in doc.tables for r in t.rows for c in r.cells)
    print(f"\nDOCX: {name}")
    print(f"  段落={len(doc.paragraphs)}  表格={ntab}  正文字符={nchars}  表格字符={tchars}")
