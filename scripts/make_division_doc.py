# -*- coding: utf-8 -*-
"""
生成交付物 №7：小组分工说明（DOCX）

包含：成员分工、技术栈选择理由、贡献说明、进度记录。
数据来源：data/verify_report.json、test/compliance_report.json（真实运行结果）
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
VERIFY = json.loads((ROOT / "data" / "verify_report.json").read_text(encoding="utf-8"))
REPORT = json.loads((ROOT / "test" / "compliance_report.json").read_text(encoding="utf-8"))
sys.stdout.reconfigure(encoding="utf-8")


def set_cjk(doc):
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return p


def para(doc, text, bold=False, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p


def bullets(doc, items, size=9.5):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet")
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.name = "微软雅黑"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9)
            r.font.name = "微软雅黑"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    return t


def main() -> int:
    doc = Document()
    set_cjk(doc)

    # ---------------- 封面 ----------------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("某课程知识库的智能问答系统\n小组分工说明")
    r.bold = True
    r.font.size = Pt(22)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("选课教材：《计算机体系结构基础》  |  技术方案：方案 C")
    r.font.size = Pt(11)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    doc.add_paragraph()
    para(doc, "说明：本文件为模板。请各组按下表结构填写本组成员真实姓名与分工，"
              "「实际完成工作」一栏建议引用本仓库中的具体文件与数据，以体现可核查的贡献。",
         size=9.5)

    # ---------------- 一、基本信息 ----------------
    h(doc, "一、小组基本信息", 1)
    table(doc, ["项目", "内容"], [
        ["课程", "计算机体系结构"],
        ["教材", "《计算机体系结构基础》胡伟武 等（430 页，12 章）"],
        ["技术栈方案", "方案 C —— SQLite FTS5 + 直接调用 LLM API + 自定义检索逻辑"],
        ["小组人数", "（请填写，要求 3~4 人）"],
        ["组长", "（请填写）"],
    ], widths=[Pt(100), Pt(340)])

    # ---------------- 二、成员分工 ----------------
    h(doc, "二、成员分工", 1)
    para(doc, "参考作业建议的四类角色，本组分工如下（请按实际填写姓名）：")
    table(doc, ["成员", "角色", "主要职责", "对应交付物"], [
        ["成员 A", "知识库构建",
         "PDF 文本提取、章节边界检测、语义分块、FTS5 建库、知识库验证",
         "data/arch_kb.sqlite\nscripts/01~04_*.py\nverify_report.json"],
        ["成员 B", "智能体配置",
         "检索工具实现、系统提示词设计、DSH 插件与智能体配置",
         "scripts/knowledge_search.py\nprompts/system_prompt.txt\nagents/*"],
        ["成员 C", "测试验证",
         "冒烟测试设计、R1~R6 自动核查脚本、合规性报告撰写",
         "test/run_smoke_test.py\ncompliance_report.json\n测试报告"],
        ["成员 D", "汇报展示",
         "PPT 制作、课堂演示、文档整理与归档",
         "实施过程 PPT\n本分工说明\nREADME"],
    ], widths=[Pt(50), Pt(70), Pt(200), Pt(120)])

    para(doc, "")
    para(doc, "填写表（请补充真实姓名与贡献比例）：", bold=True)
    table(doc, ["姓名", "学号", "角色", "贡献比例", "签名"], [
        ["", "", "成员 A 知识库构建", "", ""],
        ["", "", "成员 B 智能体配置", "", ""],
        ["", "", "成员 C 测试验证", "", ""],
        ["", "", "成员 D 汇报展示", "", ""],
    ], widths=[Pt(80), Pt(90), Pt(140), Pt(70), Pt(60)])

    # ---------------- 三、技术栈选择理由 ----------------
    doc.add_page_break()
    h(doc, "三、技术栈选择理由", 1)
    para(doc, "3.1 选定方案：方案 C（从零搭建）", bold=True)
    para(doc,
         "方案 C 的技术栈为：SQLite + FTS5（纯 Python 建库）作为知识库，"
         "直接调用 LLM API 并配合自定义 prompt 与自研检索逻辑。")

    para(doc, "3.2 与其它方案的对比", bold=True)
    table(doc, ["方案", "技术特点", "未选用的理由"], [
        ["A（参考路径）", "DSH 框架 + 自定义插件，SQLite FTS5 + BM25",
         "属教师参考路径。作业注明「选择非参考路径且实现完整者可获满分」，"
         "故选 C 以体现差异化；但仍额外提供了 DSH 插件实现以兼容 A 路径"],
        ["B", "ChromaDB/FAISS + LangChain/LlamaIndex，向量检索",
         "向量检索需引入 embedding 模型与向量库，依赖重；"
         "对于「中文教材 + 精确术语检索」场景，BM25 的确定性与可解释性更好，"
         "便于逐条验证 R1~R6"],
        ["D", "Dify/FastGPT 等低代码平台",
         "平台封装了检索过程，难以控制「先检索后回答」的执行顺序，"
         "不利于 R1 的强制与取证"],
    ], widths=[Pt(70), Pt(160), Pt(210)])

    para(doc, "3.3 选择理由小结", bold=True)
    bullets(doc, [
        "理解深度：检索链完全自研，能清晰解释每个环节的作用，符合「理解最深入」的定位。",
        "可验证性：知识库为单文件 SQLite，检索结果确定可复现，"
        "R1~R6 可被程序自动核查（本项目已实现 30 项自动核查）。",
        "依赖极简：仅需 pymupdf（建库阶段）+ Python 标准库，"
        "无需 GPU、无需向量数据库、无需框架。",
        "中文适配：选用 FTS5 trigram 分词器解决中文无空格分词问题，"
        "并内置体系结构术语同义词表。",
    ])

    # ---------------- 四、实际完成情况 ----------------
    doc.add_page_break()
    h(doc, "四、实际完成情况与数据", 1)
    para(doc, "以下数据均来自项目脚本的真实运行输出，可通过运行对应脚本复现。")

    cs = VERIFY["chunk_stats"]
    cc = VERIFY["chapter_coverage"]
    rt = VERIFY["retrieval_test"]
    per = REPORT["summary"]["per_rule"]
    total_pass = sum(v["pass"] for v in per.values())
    total_all = sum(v["pass"] + v["fail"] for v in per.values())

    para(doc, "4.1 Phase 1 知识库构建", bold=True)
    table(doc, ["指标", "结果", "复现命令"], [
        ["教材页数", "430 页", "01_extract_pages.py"],
        ["识别章节数", f"{cc['total_chapters']} 章（+6 个部分）", "01_extract_pages.py"],
        ["分块总数", f"{cs['chunk_count']} 块", "02_chunk.py"],
        ["token 均值 / 范围", f"{cs['token_avg']} / {cs['token_min']}~{cs['token_max']}",
         "02_chunk.py"],
        ["落在 800~1200", f"{cs['in_800_1200']} 块（{cs['in_800_1200_pct']}%）",
         "04_verify_kb.py"],
        ["超过上限 1200", f"{cs['above_1200']} 块", "04_verify_kb.py"],
        ["章节覆盖率", f"{cc['coverage_pct']}%（{cc['covered']}/{cc['total_chapters']}）",
         "04_verify_kb.py"],
        ["页码元数据缺失", f"{cc['bad_metadata_chunks']} 块", "04_verify_kb.py"],
        ["检索 Top-1 准确率", f"{rt['top1_accuracy']}%", "04_verify_kb.py"],
        ["检索 Top-3 命中率", f"{rt['top3_recall']}%", "04_verify_kb.py"],
    ], widths=[Pt(120), Pt(160), Pt(160)])

    para(doc, "4.2 Phase 2 智能问答系统", bold=True)
    table(doc, ["指标", "结果", "复现命令"], [
        ["检索工具", "knowledge_search（含同义词扩展、二次排序、相关性闸门）",
         "knowledge_search.py"],
        ["系统提示词", "五段式，逐段对应 R1~R6", "prompts/system_prompt.txt"],
        ["测试用例数", f"{REPORT['summary']['cases']} 个（3 正常 + 2 边缘）",
         "run_smoke_test.py"],
        ["R1~R6 核查项", f"{total_pass}/{total_all} 项通过"
                        f"（{total_pass/total_all*100:.0f}%）", "run_smoke_test.py"],
        ["合规性结论", "R1~R6 全部通过" if REPORT["summary"]["all_pass"] else "存在未通过项",
         "run_smoke_test.py"],
    ], widths=[Pt(110), Pt(200), Pt(130)])

    para(doc, "4.3 交付物清单", bold=True)
    table(doc, ["序号", "交付物", "文件", "状态"], [
        ["1", "知识库 + 构建脚本", "data/arch_kb.sqlite + scripts/01~03", "已完成"],
        ["2", "检索工具源码 + 调用说明", "scripts/knowledge_search.py + agents/README.md",
         "已完成"],
        ["3", "系统提示词 + 设计说明", "prompts/system_prompt.txt + .md", "已完成"],
        ["4", "智能体配置 + 启动说明", "agents/agent.config.json + README.md", "已完成"],
        ["5", "测试报告（DOCX）", "deliverables/Arch-Tutor_测试报告.docx", "已完成"],
        ["6", "实施过程 PPT（15~20 页）", "deliverables/Arch-Tutor_实施过程.pptx",
         "已完成"],
        ["7", "小组分工说明（DOCX）", "deliverables/Arch-Tutor_小组分工说明.docx",
         "已完成"],
    ], widths=[Pt(36), Pt(150), Pt(200), Pt(54)])

    # ---------------- 五、贡献说明 ----------------
    doc.add_page_break()
    h(doc, "五、贡献说明", 1)
    para(doc,
         "本组按「知识库构建 / 智能体配置 / 测试验证 / 汇报展示」四类角色分工，"
         "各角色工作可独立验证。建议在提交时按下列方式量化贡献：")
    table(doc, ["角色", "可验证的产出", "工作量占比（待填）"], [
        ["成员 A 知识库构建",
         "scripts/01~04；data/ 下 6 个产物文件；verify_report.json 全部指标", ""],
        ["成员 B 智能体配置",
         "knowledge_search.py（约 370 行）；system_prompt.txt；agent.config.json；"
         "arch-tutor-plugin.js", ""],
        ["成员 C 测试验证",
         "run_smoke_test.py（30 项自动核查）；compliance_report.json；测试报告 DOCX", ""],
        ["成员 D 汇报展示", "实施过程 PPT（15~20 页）；README；本分工说明", ""],
    ], widths=[Pt(110), Pt(260), Pt(70)])

    para(doc, "")
    h(doc, "5.1 协作方式", 2)
    bullets(doc, [
        "统一仓库：全部代码与文档集中在 arch-rag-kb/ 目录，按 data/scripts/prompts/"
        "agents/test/docs 分层。",
        "接口先行：知识库的产出格式（chunks.jsonl 字段：text, page_start, page_end, "
        "chapter）在动工前确定，使建库与检索可并行开发。",
        "数据驱动验收：以 verify_report.json 与 compliance_report.json 作为验收依据，"
        "避免主观判断。",
    ])

    para(doc, "")
    h(doc, "5.2 遇到的问题与解决", 2)
    table(doc, ["问题", "解决方案", "负责角色"], [
        ["中文无法用 FTS5 默认分词器检索",
         "改用 trigram 分词器，并内置体系结构术语同义词表", "成员 B"],
        ["分块出现大量 600 token 的半满块（一次 50 对）",
         "定位到 flush() 的「超限即二分」逻辑，引入 SPLIT_TOLERANCE=1.12 容忍轻微超限",
         "成员 A"],
        ["提问与教材无关时仍据无关片段作答（违反 R2/R5）",
         "新增相关性闸门 relevance_gate()，按实义关键词覆盖率（阈值 0.34）判定",
         "成员 C"],
        ["回答中出现连续 5 句原文复制（违反 R6）",
         "每片段限取 2 句、单条截断 160 字符，并过滤页眉/图注残片", "成员 B"],
        ["检索「指令流水线」未命中对应章节",
         "定位到 SQL 直接 LIMIT topk 导致候选池过小；改为先召回 topk×8 再二次排序",
         "成员 B"],
    ], widths=[Pt(150), Pt(230), Pt(60)])

    out = OUT / "Arch-Tutor_小组分工说明.docx"
    doc.save(out)
    print(f"已生成: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
